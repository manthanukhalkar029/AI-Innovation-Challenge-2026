"""
Extracts plain text (with light structure hints) from any of the required
learning-material formats: PDF, DOCX, PPTX, TXT - plus, via
`ingestion/vision.py`, OCR text pulled out of any embedded images,
diagrams, or scanned pages that native text extraction can't see
(Computer Vision, section 16 technology list).
"""
import logging
import os

logger = logging.getLogger("ai_teacher.extract")


def extract_text(filepath: str) -> str:
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".pdf":
        return _extract_pdf(filepath)
    if ext == ".docx":
        return _extract_docx(filepath)
    if ext in (".pptx",):
        return _extract_pptx(filepath)
    if ext in (".txt", ".md"):
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    raise ValueError(f"Unsupported file type: {ext}")


def _extract_pdf(filepath: str) -> str:
    import pdfplumber
    from .vision import ocr_pdf_scanned_pages

    text_parts = []
    native_text_by_page = {}
    with pdfplumber.open(filepath) as pdf:
        for i, page in enumerate(pdf.pages):
            page_text = page.extract_text() or ""
            native_text_by_page[i] = page_text
            if page_text.strip():
                text_parts.append(f"[Page {i+1}]\n{page_text}")

    try:
        ocr_parts = ocr_pdf_scanned_pages(filepath, native_text_by_page)
        text_parts.extend(ocr_parts)
    except Exception as e:
        logger.warning("PDF image/OCR pass skipped (%s: %s)", type(e).__name__, e)

    return "\n\n".join(text_parts)


def _extract_docx(filepath: str) -> str:
    import docx
    from .vision import extract_images_docx, ocr_images

    doc = docx.Document(filepath)
    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            style = (p.style.name or "").lower()
            prefix = "## " if "heading" in style else ""
            parts.append(prefix + p.text)
    for t in doc.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))

    try:
        for text in ocr_images(extract_images_docx(filepath)):
            parts.append(f"[Image - OCR]\n{text}")
    except Exception as e:
        logger.warning("DOCX image/OCR pass skipped (%s: %s)", type(e).__name__, e)

    return "\n".join(parts)


def _extract_pptx(filepath: str) -> str:
    from pptx import Presentation
    from .vision import extract_images_pptx, ocr_images

    prs = Presentation(filepath)
    parts = []
    for i, slide in enumerate(prs.slides):
        parts.append(f"## Slide {i+1}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs)
                    if line.strip():
                        parts.append(line)
            if shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))

    try:
        for text in ocr_images(extract_images_pptx(filepath)):
            parts.append(f"[Image - OCR]\n{text}")
    except Exception as e:
        logger.warning("PPTX image/OCR pass skipped (%s: %s)", type(e).__name__, e)

    return "\n".join(parts)
