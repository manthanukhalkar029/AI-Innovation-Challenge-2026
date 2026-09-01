"""
Step 1a — Document Ingestion
Parses uploaded PDF / DOCX / PPTX / TXT files into clean, chunked text
ready for embedding into the vector store.
"""

from __future__ import annotations
import os
import re
import uuid
from dataclasses import dataclass, field
from typing import List

from pypdf import PdfReader
from docx import Document as DocxDocument
from pptx import Presentation


@dataclass
class Chunk:
    """A single retrievable piece of the source material."""
    id: str
    text: str
    source: str          # original filename
    location: str        # e.g. "page 4" or "slide 12" or "paragraph 3"
    metadata: dict = field(default_factory=dict)


# --------------------------------------------------------------------------
# Per-format extraction — each returns a list of (raw_text, location) pairs
# --------------------------------------------------------------------------

def _extract_pdf(path: str) -> List[tuple[str, str]]:
    reader = PdfReader(path)
    out = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            out.append((text, f"page {i}"))
    return out


def _extract_docx(path: str) -> List[tuple[str, str]]:
    doc = DocxDocument(path)
    out = []
    buffer = []
    para_count = 0
    for para in doc.paragraphs:
        if para.text.strip():
            buffer.append(para.text)
            para_count += 1
        # flush every ~10 paragraphs to keep locations meaningful
        if len(buffer) >= 10:
            out.append((" ".join(buffer), f"paragraphs ~{para_count - 9}-{para_count}"))
            buffer = []
    if buffer:
        out.append((" ".join(buffer), f"paragraphs ~{para_count - len(buffer) + 1}-{para_count}"))
    return out


def _extract_pptx(path: str) -> List[tuple[str, str]]:
    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = "".join(run.text for run in p.runs)
                    if t.strip():
                        texts.append(t)
        if texts:
            out.append((" ".join(texts), f"slide {i}"))
    return out


def _extract_txt(path: str) -> List[tuple[str, str]]:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        content = f.read()
    # split into ~2000 char blocks so locations stay useful
    blocks = [content[i:i + 2000] for i in range(0, len(content), 2000)]
    return [(b, f"section {i+1}") for i, b in enumerate(blocks) if b.strip()]


EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".pptx": _extract_pptx,
    ".txt": _extract_txt,
    ".md": _extract_txt,
}


# --------------------------------------------------------------------------
# Cleaning + chunking
# --------------------------------------------------------------------------

def _clean(text: str) -> str:
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"(\w)-\s+(\w)", r"\1\2", text)  # de-hyphenate line breaks
    return text.strip()


def _split_into_chunks(text: str, max_chars: int = 900, overlap: int = 150) -> List[str]:
    """
    Simple sliding-window chunker on sentence boundaries.
    Good enough for Day 1 — swap for a semantic chunker later if needed.
    """
    sentences = re.split(r"(?<=[.!?])\s+", text)
    chunks, current = [], ""
    for sent in sentences:
        if len(current) + len(sent) + 1 <= max_chars:
            current = f"{current} {sent}".strip()
        else:
            if current:
                chunks.append(current)
            # start new chunk with overlap from the end of the previous one
            overlap_text = current[-overlap:] if current else ""
            current = f"{overlap_text} {sent}".strip()
    if current:
        chunks.append(current)
    return chunks


# --------------------------------------------------------------------------
# Public entry point
# --------------------------------------------------------------------------

def ingest_file(path: str) -> List[Chunk]:
    """
    Parse a file at `path` and return a list of Chunk objects
    ready to be embedded and stored.
    """
    ext = os.path.splitext(path)[1].lower()
    extractor = EXTRACTORS.get(ext)
    if extractor is None:
        raise ValueError(f"Unsupported file type: {ext}")

    filename = os.path.basename(path)
    raw_sections = extractor(path)

    chunks: List[Chunk] = []
    for raw_text, location in raw_sections:
        cleaned = _clean(raw_text)
        if not cleaned:
            continue
        for piece in _split_into_chunks(cleaned):
            chunks.append(
                Chunk(
                    id=str(uuid.uuid4()),
                    text=piece,
                    source=filename,
                    location=location,
                )
            )
    return chunks


if __name__ == "__main__":
    import sys
    if len(sys.argv) != 2:
        print("Usage: python ingestion.py <file_path>")
        sys.exit(1)
    result = ingest_file(sys.argv[1])
    print(f"Extracted {len(result)} chunks from {sys.argv[1]}")
    for c in result[:3]:
        print(f"\n--- {c.location} ---\n{c.text[:300]}...")
